import os
import json
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, numbers
)
from openpyxl.utils import get_column_letter
from core.database import (
    get_bills_for_session, get_duties_for_appointment, get_session_by_id,
    get_appointments_for_session,
)

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "outputs")

IUB_GREEN   = "1A5276"
IUB_HEADER  = "D6EAF8"
ALT_ROW     = "EBF5FB"
GOLD        = "B7950B"
WHITE       = "FFFFFF"

_thin = Side(style="thin", color="AAAAAA")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _hdr_font():  return Font(name="Calibri", bold=True, color=WHITE, size=10)
def _bold():      return Font(name="Calibri", bold=True, size=10)
def _normal():    return Font(name="Calibri", size=10)
def _center():    return Alignment(horizontal="center", vertical="center", wrap_text=True)
def _left():      return Alignment(horizontal="left",   vertical="center", wrap_text=True)
def _right():     return Alignment(horizontal="right",  vertical="center")
def _green_fill():return PatternFill("solid", fgColor=IUB_GREEN)
def _hdr_fill():  return PatternFill("solid", fgColor=IUB_HEADER)
def _alt_fill():  return PatternFill("solid", fgColor=ALT_ROW)


def _iub_header_row(ws, col_count, session_name):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=col_count)

    ws.cell(1,1).value = "The Islamia University of Bahawalpur — Examinations Branch"
    ws.cell(1,1).font  = Font(name="Calibri", bold=True, size=13, color=WHITE)
    ws.cell(1,1).fill  = _green_fill()
    ws.cell(1,1).alignment = _center()
    ws.row_dimensions[1].height = 22

    ws.cell(2,1).value = f"Session: {session_name}"
    ws.cell(2,1).font  = Font(name="Calibri", bold=True, size=10, color=IUB_GREEN)
    ws.cell(2,1).fill  = PatternFill("solid", fgColor=IUB_HEADER)
    ws.cell(2,1).alignment = _center()

    ws.cell(3,1).value = f"Generated: {datetime.now().strftime('%d-%m-%Y %H:%M')}"
    ws.cell(3,1).font  = Font(name="Calibri", size=9, color="777777")
    ws.cell(3,1).alignment = _center()
    ws.row_dimensions[3].height = 14


def _col_headers(ws, row, headers):
    for col, h in enumerate(headers, 1):
        c = ws.cell(row, col)
        c.value = h
        c.font  = _hdr_font()
        c.fill  = _green_fill()
        c.alignment = _center()
        c.border = _border
    ws.row_dimensions[row].height = 18


def _data_row(ws, row_num, values, alt=False):
    fill = _alt_fill() if alt else PatternFill("solid", fgColor=WHITE)
    for col, v in enumerate(values, 1):
        c = ws.cell(row_num, col)
        c.value = v
        c.font  = _normal()
        c.fill  = fill
        c.border = _border
        if isinstance(v, (int, float)):
            c.alignment = _right()
        else:
            c.alignment = _left()


def export_session_excel(session_id):
    session = get_session_by_id(session_id)
    if not session:
        raise ValueError(f"Session {session_id} not found.")
    session_name = session["session_name"]
    bills = get_bills_for_session(session_id)

    wb = Workbook()

    # ── Sheet 1: All Bills ────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "All Bills"
    cols1 = ["Bill No.", "Full Name", "CNIC", "Role", "Centre", "Letter No.",
             "Single Sessions", "Double Sessions", "Rate Single", "Rate Double", "Total Amount (Rs.)"]
    _iub_header_row(ws1, len(cols1), session_name)
    _col_headers(ws1, 4, cols1)
    ws1.freeze_panes = "A5"

    grand_total = 0.0
    for i, b in enumerate(bills):
        alt = (i % 2 == 1)
        row = [
            b["bill_no"], b["full_name"], b["cnic"], b["role"], b["centre"] or "",
            b["letter_no"] or "", b["total_days"], b["total_double"],
            b["rate_single"], b["rate_double"], b["amount"],
        ]
        _data_row(ws1, 5+i, row, alt)
        grand_total += b["amount"]

    total_row = 5 + len(bills)
    ws1.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=10)
    c = ws1.cell(total_row, 1)
    c.value = "GRAND TOTAL"
    c.font  = Font(name="Calibri", bold=True, size=10, color=WHITE)
    c.fill  = _green_fill()
    c.alignment = _right()
    c.border = _border
    ct = ws1.cell(total_row, 11)
    ct.value  = grand_total
    ct.font   = Font(name="Calibri", bold=True, size=10, color=WHITE)
    ct.fill   = _green_fill()
    ct.alignment = _right()
    ct.border = _border
    ct.number_format = '#,##0'

    widths1 = [18, 24, 18, 24, 24, 18, 14, 14, 12, 12, 18]
    for col, w in enumerate(widths1, 1):
        ws1.column_dimensions[get_column_letter(col)].width = w
    for row in ws1.iter_rows(min_row=5, max_row=total_row-1, min_col=9, max_col=11):
        for cell in row:
            cell.number_format = '#,##0'

    # ── Sheet 2: Summary by Role ──────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary by Role")
    cols2 = ["Role", "No. of Staff", "Total Amount (Rs.)", "% of Total"]
    _iub_header_row(ws2, len(cols2), session_name)
    _col_headers(ws2, 4, cols2)
    ws2.freeze_panes = "A5"

    role_data = {}
    for b in bills:
        r = b["role"]
        if r not in role_data:
            role_data[r] = {"count": 0, "total": 0.0}
        role_data[r]["count"] += 1
        role_data[r]["total"] += b["amount"]

    session_grand = sum(d["total"] for d in role_data.values()) or 1
    for i, (role, d) in enumerate(role_data.items()):
        pct = round(d["total"] / session_grand * 100, 1)
        _data_row(ws2, 5+i, [role, d["count"], d["total"], pct], i % 2 == 1)
        ws2.cell(5+i, 3).number_format = '#,##0'
        ws2.cell(5+i, 4).number_format = '0.0"%"'

    tr2 = 5 + len(role_data)
    ws2.merge_cells(start_row=tr2, start_column=1, end_row=tr2, end_column=2)
    c2 = ws2.cell(tr2, 1)
    c2.value = "TOTAL"
    c2.font  = Font(name="Calibri", bold=True, color=WHITE)
    c2.fill  = _green_fill()
    c2.alignment = _right()
    c2.border = _border
    ct2 = ws2.cell(tr2, 3)
    ct2.value  = sum(d["total"] for d in role_data.values())
    ct2.font   = Font(name="Calibri", bold=True, color=WHITE)
    ct2.fill   = _green_fill()
    ct2.alignment = _right()
    ct2.border = _border
    ct2.number_format = '#,##0'
    ws2.cell(tr2, 4).fill = _green_fill()
    ws2.cell(tr2, 4).border = _border

    for col, w in zip(range(1,5), [28, 14, 20, 12]):
        ws2.column_dimensions[get_column_letter(col)].width = w

    # ── Sheet 3: Duty Details ─────────────────────────────────────────────────
    ws3 = wb.create_sheet("Duty Details")
    cols3 = ["Bill No.", "Name", "Role", "Centre", "Total Sessions", "Double Sessions", "Amount Payable (Rs.)"]
    _iub_header_row(ws3, len(cols3), session_name)
    _col_headers(ws3, 4, cols3)
    ws3.freeze_panes = "A5"

    for i, b in enumerate(bills):
        _data_row(ws3, 5+i, [
            b["bill_no"], b["full_name"], b["role"], b["centre"] or "",
            b["total_days"] + b["total_double"],
            b["total_double"],
            b["amount"],
        ], i % 2 == 1)
        ws3.cell(5+i, 7).number_format = '#,##0'

    for col, w in zip(range(1,8), [18, 24, 24, 24, 14, 16, 20]):
        ws3.column_dimensions[get_column_letter(col)].width = w

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    safe_name = session_name.replace(" ", "_").replace("/", "-")
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"Session_{safe_name}_{ts}.xlsx"
    path = os.path.join(OUTPUT_DIR, filename)
    wb.save(path)
    return path


def export_staff_roster(session_id):
    session = get_session_by_id(session_id)
    if not session:
        raise ValueError(f"Session {session_id} not found.")
    session_name = session["session_name"]
    appointments = get_appointments_for_session(session_id)

    # Role priority: Superintendent=0, Deputy Superintendent=1, everything else=2
    def _role_priority(role):
        if role == "Superintendent":
            return 0
        if role == "Deputy Superintendent":
            return 1
        return 2

    # Group by centre, preserving role sort order within each centre
    centres = {}
    for a in appointments:
        centre = a["centre"] or "Unassigned"
        centres.setdefault(centre, []).append(dict(a))
    for centre in centres:
        centres[centre].sort(key=lambda x: _role_priority(x["role"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Staff Roster"

    cols = ["#", "Name", "CNIC", "Role", "Designation", "Institution", "Centre", "Letter No.", "Phone"]
    col_count = len(cols)
    _iub_header_row(ws, col_count, session_name)

    # Sheet subtitle
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=col_count)
    sub = ws.cell(4, 1)
    sub.value = "STAFF ROSTER — ALL CENTRES"
    sub.font = Font(name="Calibri", bold=True, size=11, color=IUB_GREEN)
    sub.alignment = _center()
    ws.row_dimensions[4].height = 16

    _col_headers(ws, 5, cols)
    ws.freeze_panes = "A6"

    current_row = 6
    overall_seq = 0

    for centre_name in sorted(centres.keys()):
        # Centre header band
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=col_count)
        ch = ws.cell(current_row, 1)
        ch.value = f"  Centre: {centre_name}"
        ch.font = Font(name="Calibri", bold=True, size=10, color=WHITE)
        ch.fill = PatternFill("solid", fgColor="154360")
        ch.alignment = _left()
        ch.border = _border
        ws.row_dimensions[current_row].height = 16
        current_row += 1

        invig_seq = 0
        for a in centres[centre_name]:
            overall_seq += 1
            role = a["role"]
            if role == "Invigilator":
                invig_seq += 1
                seq_label = f"Inv-{invig_seq}"
            else:
                seq_label = str(overall_seq)

            row_data = [
                seq_label,
                a.get("full_name", ""),
                a.get("cnic", ""),
                role,
                a.get("designation", "") or "",
                a.get("institution", "") or "",
                centre_name,
                a.get("letter_no", "") or "",
                a.get("phone", "") or "",
            ]
            alt = (overall_seq % 2 == 0)
            _data_row(ws, current_row, row_data, alt)

            # Bold the Superintendent and Deputy rows
            if role in ("Superintendent", "Deputy Superintendent"):
                for col in range(1, col_count + 1):
                    ws.cell(current_row, col).font = Font(name="Calibri", bold=True, size=10)

            current_row += 1

        # Blank separator row between centres
        current_row += 1

    # Grand count footer
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=col_count - 1)
    fc = ws.cell(current_row, 1)
    fc.value = f"Total Appointments: {overall_seq}"
    fc.font = Font(name="Calibri", bold=True, size=10, color=WHITE)
    fc.fill = _green_fill()
    fc.alignment = _right()
    fc.border = _border
    ws.cell(current_row, col_count).fill = _green_fill()
    ws.cell(current_row, col_count).border = _border

    col_widths = [6, 26, 18, 24, 22, 28, 24, 16, 18]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    safe_name = session_name.replace(" ", "_").replace("/", "-")
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    filename = f"Roster_{safe_name}_{ts}.xlsx"
    path = os.path.join(OUTPUT_DIR, filename)
    wb.save(path)
    return path
